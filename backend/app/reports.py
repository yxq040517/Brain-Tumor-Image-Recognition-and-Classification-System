from __future__ import annotations

from io import BytesIO
from pathlib import Path
from threading import Lock
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image as PillowImage
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFError, TTFont
from reportlab.platypus import (
    Flowable,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from .schemas import LABELS, Analysis

_LABEL_NAMES = {
    "No Tumor": "无肿瘤类",
    "Glioma": "胶质瘤类",
    "Meningioma": "脑膜瘤类",
    "Pituitary": "垂体瘤类",
}
_NAVY = colors.HexColor("#17243D")
_VIOLET = colors.HexColor("#7C3AED")
_PALE = colors.HexColor("#F3F0FF")
_BORDER = colors.HexColor("#DFE5EF")
_FONT_NAME = "BrainReportCJK"
_FONT_LOCK = Lock()
_FALLBACK_FONTS: dict[int, str | None] = {}
_FALLBACK_PATHS = (
    Path("C:/Windows/Fonts/msyh.ttc"),
    Path("C:/Windows/Fonts/simsunb.ttf"),
    Path("C:/Windows/Fonts/seguisym.ttf"),
    Path("C:/Windows/Fonts/seguiemj.ttf"),
    Path("C:/Windows/Fonts/malgun.ttf"),
)
_GLYPH_FONTS: dict[int, str] = {}


def _report_font() -> str:
    with _FONT_LOCK:
        if _FONT_NAME in pdfmetrics.getRegisteredFontNames():
            return _FONT_NAME
        # SimHei is an embeddable TrueType font; some YaHei TTC releases use
        # outlines unsupported by ReportLab, so the confirmed TTF is preferred.
        candidates = (
            Path("C:/Windows/Fonts/simhei.ttf"),
            Path("C:/Windows/Fonts/msyh.ttc"),
            Path("C:/Windows/Fonts/simsun.ttc"),
            Path("/usr/share/fonts/truetype/arphic/uming.ttc"),
            Path("/usr/share/fonts/truetype/arphic/ukai.ttc"),
            Path("/usr/share/fonts/truetype/noto/NotoSansSC-Regular.ttf"),
        )
        for path in candidates:
            if not path.is_file():
                continue
            try:
                font = TTFont(_FONT_NAME, str(path), subfontIndex=0)
                if not all(
                    ord(character) in font.face.charWidths
                    for character in "影像分析报告胶质瘤脑膜垂体备注复核"
                ):
                    continue
                pdfmetrics.registerFont(font)
                return _FONT_NAME
            except (ValueError, OSError, TTFError, pdfmetrics.FontError):
                continue
        raise RuntimeError(
            "PDF 导出需要可嵌入的中文字体，请安装 SimHei 或兼容的中文 TrueType 字体。"
        )


def _glyph_font(character: str, primary: str) -> str:
    code = ord(character)
    if character.isspace() or code in pdfmetrics.getFont(primary).face.charWidths:
        return primary
    with _FONT_LOCK:
        if code in _GLYPH_FONTS:
            return _GLYPH_FONTS[code]
        for index, path in enumerate(_FALLBACK_PATHS):
            if index not in _FALLBACK_FONTS:
                _FALLBACK_FONTS[index] = None
                if path.is_file():
                    name = f"BrainReportFallback{index}"
                    try:
                        pdfmetrics.registerFont(TTFont(name, str(path), subfontIndex=0))
                        _FALLBACK_FONTS[index] = name
                    except (ValueError, OSError, TTFError, pdfmetrics.FontError):
                        continue
            name = _FALLBACK_FONTS[index]
            if name and code in pdfmetrics.getFont(name).face.charWidths:
                _GLYPH_FONTS[code] = name
                return name
    raise ValueError(f"PDF 字体不包含字符 U+{code:04X}，请安装支持该字符的字体。")


def _font_runs(text: str, primary: str) -> list[tuple[str, str]]:
    runs: list[tuple[str, str]] = []
    current_font = primary
    characters: list[str] = []
    for character in text:
        font = _glyph_font(character, primary)
        if font != current_font and characters:
            runs.append((current_font, "".join(characters)))
            characters.clear()
        current_font = font
        characters.append(character)
    if characters:
        runs.append((current_font, "".join(characters)))
    return runs


def _paragraph(value: object, style: ParagraphStyle) -> Paragraph:
    text = str(value) if value is not None else ""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = []
    for line in text.split("\n"):
        lines.append(
            "".join(
                f'<font name="{font}">{escape(run)}</font>'
                for font, run in _font_runs(line, style.fontName)
            )
        )
    return Paragraph("<br/>".join(lines), style)


def _label(value: str | None) -> str:
    return _LABEL_NAMES.get(value, value) if value else "—"


def _wrap_label(text: str, font: str, size: float, width: float) -> list[str]:
    lines: list[str] = []
    for paragraph in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        line = ""
        line_width = 0.0
        for character in paragraph:
            character_width = pdfmetrics.stringWidth(character, _glyph_font(character, font), size)
            if line and line_width + character_width > width:
                lines.append(line)
                line = character
                line_width = character_width
            else:
                line += character
                line_width += character_width
        lines.append(line)
    return lines


class _AnnotatedImage(Flowable):
    def __init__(self, path: Path, record: Analysis, font: str) -> None:
        super().__init__()
        self.annotations = record.annotations
        self.font = font
        with PillowImage.open(path) as source:
            self.width = min(170 * mm, 92 * mm * source.width / source.height)
            self.height = self.width * source.height / source.width
            image = source.convert("RGB")
            try:
                image.thumbnail((1400, 1400), PillowImage.Resampling.LANCZOS)
                self.buffer = BytesIO()
                image.save(self.buffer, format="JPEG", quality=88, optimize=True)
            finally:
                image.close()
        self.reader = ImageReader(self.buffer)
        self.hAlign = "CENTER"

    def draw(self) -> None:
        canvas = self.canv
        canvas.saveState()
        try:
            canvas.drawImage(self.reader, 0, 0, self.width, self.height)
            canvas.setLineWidth(1.3)
            for index, annotation in enumerate(self.annotations, start=1):
                color = colors.HexColor(annotation.color)
                x = annotation.x * self.width
                y = (1 - annotation.y - annotation.height) * self.height
                width = annotation.width * self.width
                height = annotation.height * self.height
                canvas.setStrokeColor(color)
                canvas.rect(x, y, width, height, stroke=1, fill=0)
                label = f"{index}. {annotation.text}" if annotation.text else str(index)
                size = 9
                label_width = min(max(width, 70), self.width)
                lines = _wrap_label(label, self.font, size, label_width - 8)
                # Full annotation text also appears below the image and can paginate.
                max_lines = max(1, min(3, int((self.height - 6) / 12)))
                if len(lines) > max_lines:
                    lines = lines[:max_lines]
                    lines[-1] = lines[-1].rstrip() + "…"
                box_height = len(lines) * 12 + 6
                label_x = min(x, max(0, self.width - label_width))
                label_y = min(max(0, y + height - box_height), max(0, self.height - box_height))
                canvas.setFillColor(color)
                canvas.rect(label_x, label_y, label_width, box_height, stroke=0, fill=1)
                canvas.setFillColor(colors.white)
                for line_index, line in enumerate(lines):
                    text_x = label_x + 4
                    text_y = label_y + box_height - 12 - line_index * 12
                    for run_font, run in _font_runs(line, self.font):
                        canvas.setFont(run_font, size)
                        canvas.drawString(text_x, text_y, run)
                        text_x += pdfmetrics.stringWidth(run, run_font, size)
        finally:
            canvas.restoreState()


def build_pdf(analyses: list[Analysis], image_paths: dict[str, Path]) -> bytes:
    font = _report_font()
    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=21 * mm,
        bottomMargin=18 * mm,
        title="影像分析报告",
        author="脑瘤影像识别分类系统",
        pageCompression=1,
    )
    body = ParagraphStyle(
        "ReportBody",
        fontName=font,
        fontSize=10,
        leading=15,
        textColor=_NAVY,
        wordWrap="CJK",
        splitLongWords=True,
    )
    title = ParagraphStyle("ReportTitle", parent=body, fontSize=23, leading=30, spaceAfter=12)
    heading = ParagraphStyle(
        "ReportHeading",
        parent=body,
        fontSize=13,
        leading=19,
        spaceBefore=8,
        spaceAfter=7,
        textColor=_VIOLET,
    )
    value_style = ParagraphStyle("ReportValue", parent=body, fontSize=9.5, leading=14)
    centered = ParagraphStyle("ReportCentered", parent=value_style, alignment=TA_CENTER)
    story: list[Flowable] = []
    story.append(_paragraph("影像分析报告", title))
    story.append(_paragraph(f"共 {len(analyses)} 条影像记录", body))
    story.append(Spacer(1, 8 * mm))

    def table(rows: list[list[object]], widths: list[float], header: bool = False) -> Table:
        result = Table(rows, colWidths=widths, hAlign="LEFT", repeatRows=1 if header else 0)
        commands = [
            ("FONTNAME", (0, 0), (-1, -1), font),
            ("TEXTCOLOR", (0, 0), (-1, -1), _NAVY),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 10),
            ("RIGHTPADDING", (0, 0), (-1, -1), 10),
            ("TOPPADDING", (0, 0), (-1, -1), 7),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ("LINEBELOW", (0, 0), (-1, -1), 0.4, _BORDER),
        ]
        if header:
            commands.append(("BACKGROUND", (0, 0), (-1, 0), _PALE))
        else:
            commands.append(("BACKGROUND", (0, 0), (0, -1), _PALE))
        result.setStyle(TableStyle(commands))
        return result

    width = document.width
    for index, record in enumerate(analyses):
        if index:
            story.append(PageBreak())
        story.append(_paragraph(f"{index + 1:02d}  {record.filename}", heading))
        metadata = (
            ("影像名称", record.filename),
            ("分组", record.group or "未分组"),
            ("分析时间", record.created_at),
            (
                "尺寸 / 推理耗时",
                f"{record.width} × {record.height} px / {record.inference_ms:.1f} ms",
            ),
            ("原始分类", f"{_label(record.predicted_label)} · {record.confidence:.2%}"),
            ("复核状态", "已复核" if record.review_status == "reviewed" else "待复核"),
            ("复核分类 / 时间", f"{_label(record.reviewed_label)} / {record.reviewed_at or '—'}"),
        )
        story.append(
            table(
                [
                    [_paragraph(key, value_style), _paragraph(value, value_style)]
                    for key, value in metadata
                ],
                [37 * mm, width - 37 * mm],
            )
        )
        story.append(Spacer(1, 6 * mm))
        if record.id not in image_paths:
            raise FileNotFoundError(f"影像记录 {record.id} 的源文件不存在。")
        story.append(_AnnotatedImage(image_paths[record.id], record, font))
        story.append(Spacer(1, 5 * mm))
        scores = {probability.label: probability.score for probability in record.probabilities}
        score_rows = [
            [_paragraph(_label(label), centered) for label in LABELS],
            [_paragraph(f"{scores[label]:.2%}", centered) for label in LABELS],
        ]
        story.append(table(score_rows, [width / len(LABELS)] * len(LABELS), header=True))
        story.append(_paragraph("备注", heading))
        story.append(_paragraph(record.notes or "—", body))
        if record.annotations:
            story.append(_paragraph("保存的标注", heading))
            for annotation_index, annotation in enumerate(record.annotations, start=1):
                story.append(
                    _paragraph(f"{annotation_index}. {annotation.text or '矩形标注'}", body)
                )
        story.append(_paragraph("模型版本", heading))
        story.append(
            _paragraph(
                f"{record.model_metadata.get('name') or record.model_id}\n"
                f"{record.model_id}\n{record.model_revision}",
                value_style,
            )
        )
        story.append(_paragraph(f"记录编号：{record.id}", value_style))

    def decorate(canvas, doc) -> None:
        canvas.saveState()
        canvas.setStrokeColor(_BORDER)
        canvas.line(18 * mm, A4[1] - 15 * mm, A4[0] - 18 * mm, A4[1] - 15 * mm)
        canvas.setFont(font, 9)
        canvas.setFillColor(_NAVY)
        canvas.drawString(18 * mm, A4[1] - 12 * mm, "影像分析报告")
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"第 {doc.page} 页")
        canvas.restoreState()

    document.build(story, onFirstPage=decorate, onLaterPages=decorate)
    return buffer.getvalue()


def _excel_text(value: object) -> str:
    text = str(value) if value is not None else ""
    if len(text) > 32767:
        raise ValueError("记录文本超过 Excel 单元格的 32767 字符限制，无法无损导出。")
    return text


def build_excel(analyses: list[Analysis]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "影像分析记录"
    columns = [
        ("记录编号", 39),
        ("影像名称", 32),
        ("分组", 22),
        ("分析时间", 29),
        ("宽度 (px)", 13),
        ("高度 (px)", 13),
        ("原始分类", 17),
        ("原始置信度", 16),
        ("无肿瘤类概率", 17),
        ("胶质瘤类概率", 17),
        ("脑膜瘤类概率", 17),
        ("垂体瘤类概率", 17),
        ("推理耗时 (ms)", 18),
        ("复核状态", 15),
        ("复核分类", 17),
        ("复核时间", 29),
        ("备注", 60),
        ("模型名称", 32),
        ("模型标识", 44),
        ("模型版本", 44),
        ("标注数量", 13),
    ]
    last_column = get_column_letter(len(columns))
    sheet.merge_cells(f"A1:{last_column}1")
    sheet["A1"] = "影像分析报告"
    sheet["A1"].font = Font(name="Microsoft YaHei", size=20, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="17243D")
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 42
    sheet.row_dimensions[2].height = 29
    for index, (name, width) in enumerate(columns, start=1):
        cell = sheet.cell(2, index, name)
        cell.font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="7C3AED")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = width
    thin = Side(style="hair", color="DFE5EF")
    body_font = Font(name="Microsoft YaHei", size=11, color="17243D")
    for row_number, record in enumerate(analyses, start=3):
        scores = {probability.label: probability.score for probability in record.probabilities}
        values = [
            record.id,
            record.filename,
            record.group,
            record.created_at,
            record.width,
            record.height,
            _label(record.predicted_label),
            record.confidence,
            *(scores[label] for label in LABELS),
            record.inference_ms,
            "已复核" if record.review_status == "reviewed" else "待复核",
            _label(record.reviewed_label) if record.reviewed_label else "",
            record.reviewed_at or "",
            record.notes,
            record.model_metadata.get("name") or record.model_id,
            record.model_id,
            record.model_revision,
            len(record.annotations),
        ]
        for column_number, value in enumerate(values, start=1):
            cell = sheet.cell(row_number, column_number)
            if isinstance(value, str):
                cell.value = _excel_text(value)
                # Set the actual cell type, preserving leading = + - @ literally.
                # Prefixing an apostrophe would alter the user's stored text.
                cell.data_type = "s"
            else:
                cell.value = value
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
            if row_number % 2:
                cell.fill = PatternFill("solid", fgColor="F4F6FC")
            if 8 <= column_number <= 12:
                cell.number_format = "0.00%"
            elif column_number == 13:
                cell.number_format = "0.0"
        sheet.row_dimensions[row_number].height = None
    sheet.freeze_panes = "D3"
    sheet.auto_filter.ref = f"A2:{last_column}{max(2, sheet.max_row)}"
    sheet.sheet_view.showGridLines = False
    sheet.print_title_rows = "1:2"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_options.horizontalCentered = True
    sheet.print_area = f"A1:{last_column}{max(2, sheet.max_row)}"
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()
